"""V1.1 API contracts and data safety checks. No production database access."""

from __future__ import annotations

import io
import json
import zipfile
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import inspect

from gyrotron import create_app
from gyrotron.field_definitions import FIXED_FIELDS
from gyrotron.models import CustomParameter, Shot


def client():
    app = create_app({"TESTING": True, "SQLALCHEMY_DATABASE_URI": "sqlite:///:memory:"})
    return app, app.test_client()


def make_shot(serial="TUBE-A", number="1", date="2026-10-15", iplen="1000", meaning=None):
    return {
        "Serial": serial,
        "shot": number,
        "date": date,
        "remark": None,
        "iplen": iplen,
        "Uc": "15.125",
        "Ua": None,
        "Ib": None,
        "Im": None,
        "Pf": None,
        "Ic": None,
        "Pw": None,
        "Pc": None,
        "custom_parameters": [
            {
                "name": "extra",
                "description": meaning,
                "unit": None,
                "data_type": "DECIMAL",
                "value": None,
            }
        ],
    }


def seed(test_client):
    ids = []
    for serial, number, dt, iplen, value in [
        ("TUBE-A", "1", "2026-10-13", "800000", "2.50"),
        ("TUBE-A", "2", "2026-10-14", None, None),
        ("TUBE-A", "3", "2026-10-15", "30", "5.50"),
        ("TUBE-B", "1", "2026-10-15", "-1", "8.00"),
    ]:
        row = make_shot(serial, number, dt, iplen)
        row["custom_parameters"][0]["value"] = value
        result = test_client.post("/api/v1/shots", json=row)
        assert result.status_code == 201, result.json
        ids.append(result.json["id"])
    return ids


def test_fixed_data_dictionary_and_db_types():
    assert [x[0] for x in FIXED_FIELDS] == ["iplen", "Uc", "Ua", "Ib", "Im", "Pf", "Ic", "Pw", "Pc"]
    assert [x[2] for x in FIXED_FIELDS] == ["ms", "kV", "kV", "A", "A", "W", "A", "kW", "kW"]
    names = {col.name: col for col in inspect(Shot).columns}
    assert {name: str(names[name].type) for name, _, _ in FIXED_FIELDS} == {
        name: "NUMERIC(28, 10)" for name, _, _ in FIXED_FIELDS
    }
    assert str(names["shot"].type) == "BIGINT"
    assert str(names["date"].type) == "DATE"
    assert str(names["remark"].type) == "TEXT"
    assert names["date"].nullable is False
    assert inspect(CustomParameter).columns["description"].nullable is True


def test_nullable_description_unit_value_roundtrip_and_edit():
    _, c = client()
    item = c.post("/api/v1/shots", json=make_shot())
    assert item.status_code == 201, item.json
    parameter = item.json["custom_parameters"][0]
    assert parameter["description"] is None
    assert parameter["unit"] is None
    assert parameter["value"] is None
    changed = make_shot()
    changed["custom_parameters"][0]["id"] = parameter["id"]
    changed["custom_parameters"][0]["description"] = "参数含义"
    changed["custom_parameters"][0]["value"] = "0"
    saved = c.put("/api/v1/shots/" + str(item.json["id"]), json=changed)
    assert saved.status_code == 200, saved.json
    assert saved.json["custom_parameters"][0]["description"] == "参数含义"
    assert saved.json["custom_parameters"][0]["value"] == "0"
    assert saved.json["custom_parameters"][0]["unit"] is None


def test_date_required_and_identity_immutable():
    _, c = client()
    bad = make_shot()
    bad["date"] = None
    assert c.post("/api/v1/shots", json=bad).status_code == 400
    created = c.post("/api/v1/shots", json=make_shot()).json
    record_id = created["id"]
    for key, value in (("shot", "2"), ("Serial", "TUBE-X")):
        modified = make_shot()
        modified[key] = value
        denied = c.put("/api/v1/shots/" + str(record_id), json=modified)
        assert denied.status_code == 400, denied.json
    assert c.get("/api/v1/shots/" + str(record_id)).json["shot"] == "1"
    assert c.get("/api/v1/shots/" + str(record_id)).json["Serial"] == "TUBE-A"


def test_cannot_change_fixed_definition_even_via_raw_api():
    _, c = client()
    row = make_shot()
    row["fixed_fields"] = [{"name": "Uc", "unit": "V"}]
    assert c.post("/api/v1/shots", json=row).status_code == 400
    row.pop("fixed_fields")
    row["custom_parameters"][0]["name"] = "Pc"
    assert c.post("/api/v1/shots", json=row).status_code == 400
    row["custom_parameters"][0]["name"] = "extra"
    assert c.post("/api/v1/shots", json=row).status_code == 201


def test_serial_suggestions_next_shot_and_all_history_dashboard():
    _, c = client()
    ids = seed(c)
    serials = c.get("/api/v1/serials").json
    assert [s["name"] for s in serials] == ["TUBE-A", "TUBE-B"]
    assert c.get("/api/v1/next-shot?serial=TUBE-A").json["shot"] == "4"
    suggestions = c.get("/api/v1/parameter-suggestions?serial=TUBE-A").json
    assert suggestions[0]["name"] == "extra"
    assert suggestions[0]["source"] == "TUBE-A"
    global_view = c.get("/api/v1/dashboard").json
    assert global_view["shots"] == 4
    assert global_view["serials"] == 2
    assert global_view["missing_iplen"] == 2
    assert len(global_view["points"]) == 2
    assert global_view["by_serial"] == [
        {"name": "TUBE-A", "count": 3},
        {"name": "TUBE-B", "count": 1},
    ]
    filtered = c.get("/api/v1/dashboard?serial=TUBE-A").json
    assert filtered["shots"] == 3
    assert filtered["serials"] == 2
    assert filtered["missing_iplen"] == 1
    assert filtered["latest_date"] == "2026-10-15"
    assert [x["id"] for x in filtered["points"]] == [ids[0], ids[2]]


def test_search_filters_and_full_custom_values():
    _, c = client()
    seed(c)

    def results(query):
        response = c.get("/api/v1/shots", query_string=query)
        assert response.status_code == 200, response.json
        return response.json

    r = results({"serial": "TUBE-A", "page_size": "2"})
    assert r["total"] == 3
    assert [p["shot"] for p in r["items"]] == ["3", "2"]
    assert r["items"][0]["custom_parameters"][0]["name"] == "extra"
    assert results({"date_from": "2026-10-15"})["total"] == 2
    assert results({"iplen_state": "missing"})["total"] == 2
    assert results({"iplen_min": "100"})["total"] == 1
    assert (
        results(
            {
                "custom_filters": json.dumps(
                    [{"name": "extra", "operator": "contains", "value": "5."}]
                )
            }
        )["total"]
        == 1
    )
    assert (
        results(
            {"custom_filters": json.dumps([{"name": "extra", "operator": "gt", "value": "4"}])}
        )["total"]
        == 2
    )
    assert c.get("/api/v1/shots?date_from=wat").status_code == 400
    assert c.get("/api/v1/shots?custom_filters=malformed").status_code == 400
    assert c.get("/api/v1/shots?custom_filters=[]&iplen_state=unknown").status_code == 400


def test_xlsx_and_csv_zip_all_filtered_rows():
    _, c = client()
    seed(c)
    response = c.get("/api/v1/shots/export?serial=TUBE-A&format=xlsx")
    assert response.status_code == 200, response.get_data(as_text=False)[:2]
    wb = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
    assert wb.sheetnames == ["炮次记录", "自定义参数"]
    primary = list(wb["炮次记录"].values)
    custom = list(wb["自定义参数"].values)
    assert len(primary) == 4 and len(custom) == 4
    assert "iplen (ms)" in primary[0]
    assert primary[1][1] == "1"  # shot stays a string (BIGINT precision)
    assert custom[0][3] == "参数含义"
    assert custom[1][3] is None
    zip_result = c.get("/api/v1/shots/export?serial=TUBE-A&format=csv")
    with zipfile.ZipFile(io.BytesIO(zip_result.data)) as archive:
        assert set(archive.namelist()) == {"shots.csv", "custom_parameters.csv"}
        contents = archive.read("shots.csv").decode("utf-8-sig")
        assert "TUBE-A" in contents and "TUBE-B" not in contents


def test_preview_bulk_delete_snapshot_validation_and_cascade():
    app, c = client()
    ids = seed(c)
    filters = {"serial": "TUBE-A"}
    preview = c.post("/api/v1/shots/delete-preview", json={"filters": filters})
    assert preview.status_code == 200, preview.json
    assert preview.json["count"] == 3
    # Selection changed after preview: protect unrelated records.
    new = c.post("/api/v1/shots", json=make_shot(number="4"))
    assert new.status_code == 201, new.json
    failed = c.post(
        "/api/v1/shots/bulk-delete", json={"filters": filters, "snapshot": preview.json["snapshot"]}
    )
    assert failed.status_code == 409
    assert c.get("/api/v1/shots?serial=TUBE-A").json["total"] == 4
    updated = c.post("/api/v1/shots/delete-preview", json={"filters": filters}).json
    gone = c.post(
        "/api/v1/shots/bulk-delete", json={"filters": filters, "snapshot": updated["snapshot"]}
    )
    assert gone.status_code == 200, gone.json
    assert gone.json["deleted"] == 4
    with app.app_context():
        from gyrotron.webapp import db

        assert db.session.query(Shot).count() == 1
        assert db.session.query(CustomParameter).count() == 1
    assert c.get("/api/v1/shots/" + str(ids[-1])).status_code == 200


def test_bulk_delete_selected_ids_only():
    _, c = client()
    ids = seed(c)
    filters = {"ids": [ids[0], ids[3]]}
    first = c.post("/api/v1/shots/delete-preview", json={"filters": filters})
    assert first.status_code == 200, first.json
    assert first.json["count"] == 2
    done = c.post(
        "/api/v1/shots/bulk-delete", json={"filters": filters, "snapshot": first.json["snapshot"]}
    )
    assert done.status_code == 200, done.json
    assert c.get("/api/v1/shots").json["total"] == 2


def test_type_exact_numeric_keeps_zero_separate_from_null():
    _, c = client()
    item = c.post("/api/v1/shots", json=make_shot(iplen="0"))
    assert item.status_code == 201
    assert Decimal(item.json["iplen"]) == 0
    other = c.post("/api/v1/shots", json=make_shot(number="2", iplen=None))
    assert other.status_code == 201
    assert other.json["iplen"] is None
